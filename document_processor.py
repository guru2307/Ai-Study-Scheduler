"""
document_processor.py - extract plain text from uploaded PDF / DOCX / TXT files.
"""
import os

class DocumentProcessingError(Exception):
    pass

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

def allowed_file(filename):
    ext = os.path.splitext(filename)[1].lower()
    return ext in ALLOWED_EXTENSIONS

def extract_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return _extract_pdf(filepath)
    elif ext == ".docx":
        return _extract_docx(filepath)
    elif ext == ".txt":
        return _extract_txt(filepath)
    else:
        raise DocumentProcessingError(f"Unsupported file type: {ext}")

def _extract_pdf(filepath):
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader
        except ImportError as e:
            raise DocumentProcessingError("PDF library not installed on the server.") from e

    try:
        reader = PdfReader(filepath)
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception:
                raise DocumentProcessingError("This PDF is password-protected and cannot be read.")
        text_parts = []
        for page in reader.pages:
            try:
                text_parts.append(page.extract_text() or "")
            except Exception:
                continue
        text = "\n".join(text_parts).strip()
        if not text:
            raise DocumentProcessingError(
                "No text could be extracted from this PDF. It may be a scanned image without selectable text."
            )
        return text
    except DocumentProcessingError:
        raise
    except Exception as e:
        raise DocumentProcessingError(f"Could not read this PDF file: {e}")

def _extract_docx(filepath):
    try:
        import docx
    except ImportError as e:
        raise DocumentProcessingError("DOCX library not installed on the server.") from e

    try:
        document = docx.Document(filepath)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        text = "\n".join(paragraphs).strip()
        if not text:
            raise DocumentProcessingError("This DOCX file appears to be empty.")
        return text
    except DocumentProcessingError:
        raise
    except Exception as e:
        raise DocumentProcessingError(f"Could not read this DOCX file: {e}")

def _extract_txt(filepath):
    try:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read().strip()
        if not text:
            raise DocumentProcessingError("This TXT file appears to be empty.")
        return text
    except DocumentProcessingError:
        raise
    except Exception as e:
        raise DocumentProcessingError(f"Could not read this TXT file: {e}")
